import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from app.infrastructure.extraction import ProcessExtractor, extract_document
from app.service.errors import AppError
from tests.conftest import SOURCE, docx_bytes


def pdf_bytes(*, pages=1, text=True, password=None):
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=595, height=842)
        if text:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
            stream = DecodedStreamObject()
            stream.set_data(
                b"BT /F1 12 Tf 40 750 Td (Ana Silva projeto Python PostgreSQL "
                b"2024 desenvolvimento de API para biblioteca) Tj ET"
            )
            page[NameObject("/Contents")] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_digital_pdf_and_process_isolation(settings):
    text = ProcessExtractor(settings).extract(pdf_bytes(), "file.pdf")
    assert "Python" in text and "2024" in text


def test_docx_tables_keep_block_order(settings):
    document = Document()
    document.add_paragraph("Ana Silva: estudante de programação e desenvolvimento de software.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "2024"
    document.add_paragraph("PostgreSQL no projeto acadêmico de biblioteca.")
    output = io.BytesIO()
    document.save(output)
    text = ProcessExtractor(settings).extract(output.getvalue(), "resume.docx")
    assert text.index("Ana Silva") < text.index("Python | 2024") < text.index("PostgreSQL")


@pytest.mark.parametrize(
    ("data", "filename", "code"),
    [
        (b"", "a.pdf", "arquivo_vazio"),
        (b"arbitrary", "a.exe", "formato_nao_aceito"),
        (b"not pdf", "a.pdf", "assinatura_invalida"),
        (b"%PDF-1.7\ninvalid", "a.pdf", "arquivo_corrompido"),
        (pdf_bytes(text=False), "a.pdf", "texto_insuficiente"),
        (pdf_bytes(password="secret"), "a.pdf", "pdf_protegido"),
        (pdf_bytes(pages=11), "a.pdf", "paginas_excedidas"),
    ],
)
def test_invalid_uploads(settings, data, filename, code):
    with pytest.raises(AppError) as caught:
        ProcessExtractor(settings).extract(data, filename)
    assert caught.value.code == code


def test_docx_decompression_and_text_bounds(settings):
    settings.max_docx_unpacked_bytes = 30
    with pytest.raises(AppError, match="descompactado"):
        extract_document(docx_bytes(), ".docx", settings)
    settings.max_docx_unpacked_bytes = 100_000_000
    settings.max_extracted_chars = 100
    with pytest.raises(AppError, match="Texto extraído"):
        extract_document(docx_bytes(SOURCE * 5), ".docx", settings)


def test_zip_disguised_as_docx(settings):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("file.txt", "unrelated content")
    with pytest.raises(AppError, match="DOCX inválido"):
        extract_document(output.getvalue(), ".docx", settings)


def test_upload_size_is_limited_before_parsing(settings):
    settings.max_upload_bytes = 3
    with pytest.raises(AppError) as caught:
        ProcessExtractor(settings).extract(b"1234", "a.pdf")
    assert caught.value.status == 413
